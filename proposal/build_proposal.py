"""Membangun proposal GTNIC 2026 menjadi DOCX (sumber yang bisa diedit) dan PDF.

Format mengikuti ketentuan lomba: A4, Times New Roman 12 pt, spasi 1,5,
margin kiri 4 cm serta atas/kanan/bawah 3 cm, maksimal 15 halaman.

Jalankan: C:\\Python313\\python.exe build_proposal.py
"""

from pathlib import Path

from PIL import Image as PILImage

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
                                NextPageTemplate, PageBreak, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)

import content

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "out"
OUT.mkdir(exist_ok=True)

WIN_FONTS = Path(r"C:\Windows\Fonts")
USABLE_CM = 21.0 - 4.0 - 3.0  # 14 cm
BODY_PT = 12
LEADING = BODY_PT * 1.5

pdfmetrics.registerFont(TTFont("TNR", str(WIN_FONTS / "times.ttf")))
pdfmetrics.registerFont(TTFont("TNR-Bold", str(WIN_FONTS / "timesbd.ttf")))
pdfmetrics.registerFont(TTFont("TNR-Italic", str(WIN_FONTS / "timesi.ttf")))
pdfmetrics.registerFont(TTFont("TNR-BoldItalic", str(WIN_FONTS / "timesbi.ttf")))
pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-Bold",
                              italic="TNR-Italic", boldItalic="TNR-BoldItalic")


def img_height_cm(name: str, width_cm: float) -> float:
    with PILImage.open(ASSETS / name) as im:
        w, h = im.size
    return width_cm * h / w


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

S = {
    "body": ParagraphStyle("body", fontName="TNR", fontSize=BODY_PT, leading=LEADING,
                           alignment=TA_JUSTIFY, spaceAfter=6),
    "h1": ParagraphStyle("h1", fontName="TNR-Bold", fontSize=13, leading=19,
                         spaceBefore=12, spaceAfter=5, keepWithNext=1),
    "h2": ParagraphStyle("h2", fontName="TNR-Bold", fontSize=12, leading=LEADING,
                         spaceBefore=8, spaceAfter=3, keepWithNext=1),
    "list": ParagraphStyle("list", fontName="TNR", fontSize=BODY_PT, leading=LEADING,
                           alignment=TA_JUSTIFY, leftIndent=18, bulletIndent=4,
                           spaceAfter=4),
    "cap": ParagraphStyle("cap", fontName="TNR-Italic", fontSize=11, leading=15,
                          alignment=TA_CENTER, spaceBefore=4, spaceAfter=10),
    "note": ParagraphStyle("note", fontName="TNR-Italic", fontSize=11, leading=16,
                           alignment=TA_JUSTIFY, leftIndent=8, rightIndent=8,
                           spaceBefore=4, spaceAfter=8,
                           textColor=colors.HexColor("#6b5a1f")),
    "cell": ParagraphStyle("cell", fontName="TNR", fontSize=10.5, leading=14),
    "cellh": ParagraphStyle("cellh", fontName="TNR-Bold", fontSize=10.5, leading=14),
    "cv_small": ParagraphStyle("cv_small", fontName="TNR", fontSize=12, leading=18,
                               alignment=TA_CENTER),
    "cv_mid": ParagraphStyle("cv_mid", fontName="TNR-Bold", fontSize=13, leading=20,
                             alignment=TA_CENTER),
    "cv_title": ParagraphStyle("cv_title", fontName="TNR-Bold", fontSize=18, leading=26,
                               alignment=TA_CENTER),
    "cv_logo": ParagraphStyle("cv_logo", fontName="TNR-Italic", fontSize=12, leading=18,
                              alignment=TA_CENTER,
                              textColor=colors.HexColor("#8a6d1d")),
}

TABLE_STYLE = TableStyle([
    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#5a665e")),
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfeee5")),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LEFTPADDING", (0, 0), (-1, -1), 5),
    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ("TOPPADDING", (0, 0), (-1, -1), 4),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
])

COL_WIDTHS = {3: [2.7 * cm, 5.3 * cm, 6.0 * cm]}


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("TNR", 10)
    canvas.setFillColor(colors.HexColor("#555f57"))
    canvas.drawRightString(A4[0] - 3 * cm, 1.9 * cm, str(doc.page))
    canvas.restoreState()


def pdf_cover():
    m = content.META
    flow = [Spacer(1, 0.8 * cm), Paragraph(m["logo_note"], S["cv_logo"]),
            Spacer(1, 0.9 * cm),
            Paragraph("PROPOSAL KARYA", S["cv_mid"]),
            Paragraph(m["competition"], S["cv_mid"]),
            Paragraph(m["category"], S["cv_small"]),
            Paragraph(m["subtheme"], S["cv_small"]),
            Spacer(1, 1.5 * cm),
            Paragraph(m["title"], S["cv_title"]),
            Spacer(1, 1.4 * cm),
            Paragraph(f"Nama Tim: <b>{m['team']}</b>", S["cv_small"]),
            Spacer(1, 0.3 * cm)]
    for member in m["members"]:
        flow.append(Paragraph(member, S["cv_small"]))
    flow.append(Spacer(1, 1.2 * cm))
    rows = [[Paragraph(label, S["cellh"]), Paragraph(value, S["cell"])]
            for label, value in m["links"]]
    table = Table(rows, colWidths=[3.6 * cm, 10.4 * cm], hAlign="CENTER")
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#5a665e")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    flow += [table, Spacer(1, 1.0 * cm),
             Paragraph(m["footer_note"], S["cv_logo"]),
             NextPageTemplate("body"), PageBreak()]
    return flow


def pdf_blocks(blocks):
    flow = []
    counter = 0
    for b in blocks:
        kind = b["t"]
        if kind == "h1":
            counter = 0
            flow.append(Paragraph(b["x"], S["h1"]))
        elif kind == "h2":
            counter = 0
            flow.append(Paragraph(b["x"], S["h2"]))
        elif kind == "p":
            counter = 0
            flow.append(Paragraph(b["x"], S["body"]))
        elif kind == "b":
            counter = 0
            flow.append(Paragraph(b["x"], S["list"], bulletText="\u2022"))
        elif kind == "n":
            counter += 1
            flow.append(Paragraph(b["x"], S["list"], bulletText=f"{counter}."))
        elif kind == "placeholder":
            counter = 0
            flow.append(Paragraph(b["x"], S["note"]))
        elif kind == "img":
            counter = 0
            w = min(b["w"], USABLE_CM)
            pic = Image(str(ASSETS / b["f"]), width=w * cm,
                        height=img_height_cm(b["f"], w) * cm)
            pic.hAlign = "CENTER"
            flow.append(KeepTogether([Spacer(1, 0.2 * cm), pic,
                                      Paragraph(b["cap"], S["cap"])]))
        elif kind == "table":
            counter = 0
            rows = [[Paragraph(c, S["cellh"]) for c in b["head"]]]
            rows += [[Paragraph(c, S["cell"]) for c in r] for r in b["rows"]]
            widths = COL_WIDTHS.get(len(b["head"]))
            table = Table(rows, colWidths=widths, hAlign="CENTER", repeatRows=1)
            table.setStyle(TABLE_STYLE)
            flow += [Spacer(1, 0.15 * cm), table,
                     Paragraph(b["cap"], S["cap"])]
        else:
            raise ValueError(f"blok tidak dikenal: {kind}")
    return flow


def build_pdf(path: Path, flow):
    doc = BaseDocTemplate(str(path), pagesize=A4,
                          leftMargin=4 * cm, rightMargin=3 * cm,
                          topMargin=3 * cm, bottomMargin=3 * cm,
                          title=content.META["title"],
                          author=content.META["team"])
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([
        PageTemplate(id="cover", frames=[frame]),
        PageTemplate(id="body", frames=[frame], onPage=footer),
    ])
    doc.build(flow)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def docx_setup(document):
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Heading 1", "Heading 2", "Caption"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0x1A, 0x20, 0x1D)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Heading 2"].font.size = Pt(12)
    document.styles["Caption"].font.size = Pt(11)
    document.styles["Caption"].font.italic = True
    return section


def docx_page_number_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def docx_para(document, text, *, style=None, align=None, bold=False, italic=False,
              size=None, spacing=1.5, space_after=6, indent=None):
    p = document.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    if size:
        run.font.size = Pt(size)
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def docx_table(document, head, rows):
    table = document.add_table(rows=1, cols=len(head))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, head):
        cell.text = ""
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.name = "Times New Roman"
        cell.paragraphs[0].paragraph_format.line_spacing = 1
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(10.5)
            run.font.name = "Times New Roman"
            cell.paragraphs[0].paragraph_format.line_spacing = 1
    return table


def build_docx(path: Path):
    m = content.META
    document = docx.Document()
    section = docx_setup(document)

    center = WD_ALIGN_PARAGRAPH.CENTER
    docx_para(document, m["logo_note"], align=center, italic=True, space_after=18)
    docx_para(document, "PROPOSAL KARYA", align=center, bold=True, size=13, space_after=0)
    docx_para(document, m["competition"], align=center, bold=True, size=13, space_after=0)
    docx_para(document, m["category"], align=center, space_after=0)
    docx_para(document, m["subtheme"], align=center, space_after=30)
    docx_para(document, m["title"], align=center, bold=True, size=18, space_after=36)
    docx_para(document, f"Nama Tim: {m['team']}", align=center, bold=True, space_after=6)
    for member in m["members"]:
        docx_para(document, member, align=center, space_after=0)
    docx_para(document, "", align=center, space_after=18)
    docx_table(document, ["Keterangan", "Tautan"], [list(x) for x in m["links"]])
    docx_para(document, "", space_after=12)
    docx_para(document, m["footer_note"], align=center, italic=True)

    document.add_section(WD_SECTION.NEW_PAGE)
    body = document.sections[1]
    body.page_width, body.page_height = Cm(21), Cm(29.7)
    body.left_margin, body.right_margin = Cm(4), Cm(3)
    body.top_margin, body.bottom_margin = Cm(3), Cm(3)
    body.footer.is_linked_to_previous = False
    docx_page_number_footer(body)

    counter = 0
    for b in content.BLOCKS:
        kind = b["t"]
        if kind == "h1":
            counter = 0
            document.add_heading(b["x"], level=1)
        elif kind == "h2":
            counter = 0
            document.add_heading(b["x"], level=2)
        elif kind == "p":
            counter = 0
            docx_para(document, b["x"])
        elif kind == "b":
            counter = 0
            docx_para(document, b["x"], style="List Bullet", space_after=4)
        elif kind == "n":
            counter += 1
            docx_para(document, b["x"], style="List Number", space_after=4)
        elif kind == "placeholder":
            counter = 0
            docx_para(document, b["x"], italic=True, size=11, indent=0.5)
        elif kind == "img":
            counter = 0
            w = min(b["w"], USABLE_CM)
            p = document.add_paragraph()
            p.alignment = center
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(str(ASSETS / b["f"]), width=Cm(w))
            cap = document.add_paragraph(b["cap"], style="Caption")
            cap.alignment = center
            cap.paragraph_format.line_spacing = 1
            cap.paragraph_format.space_after = Pt(10)
        elif kind == "table":
            counter = 0
            docx_table(document, b["head"], b["rows"])
            cap = document.add_paragraph(b["cap"], style="Caption")
            cap.alignment = center
            cap.paragraph_format.line_spacing = 1
            cap.paragraph_format.space_after = Pt(10)

    document.save(str(path))
    return section


# --------------------------------------------------------------------------- #
# Surat pernyataan keaslian karya (dokumen terpisah)
# --------------------------------------------------------------------------- #

def pledge_blocks():
    p = content.PLEDGE
    flow = [Paragraph(f"<b>{p['title']}</b>",
                      ParagraphStyle("pt", parent=S["cv_title"], fontSize=14, leading=22)),
            Spacer(1, 0.8 * cm), Paragraph(p["intro"], S["body"]), Spacer(1, 0.3 * cm)]
    rows = [[Paragraph(f"<b>{label}</b>", S["cell"]), Paragraph(value, S["cell"])]
            for label, value in p["fields"]]
    table = Table(rows, colWidths=[4.2 * cm, 9.8 * cm], hAlign="LEFT")
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    flow += [table, Spacer(1, 0.5 * cm),
             Paragraph("Dengan ini menyatakan bahwa:", S["body"])]
    for i, statement in enumerate(p["statements"], 1):
        flow.append(Paragraph(statement, S["list"], bulletText=f"{i}."))
    flow += [Spacer(1, 0.4 * cm), Paragraph(p["closing"], S["body"]),
             Spacer(1, 0.8 * cm)]
    right = ParagraphStyle("right", parent=S["body"], alignment=2)
    for sline in p["signature"]:
        flow.append(Paragraph(sline or "&nbsp;", right))
    return flow


def build_pledge_docx(path: Path):
    p = content.PLEDGE
    document = docx.Document()
    docx_setup(document)
    docx_para(document, p["title"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
              size=14, space_after=18)
    docx_para(document, p["intro"])
    docx_table(document, ["Keterangan", "Isi"], [list(x) for x in p["fields"]])
    docx_para(document, "", space_after=6)
    docx_para(document, "Dengan ini menyatakan bahwa:")
    for statement in p["statements"]:
        docx_para(document, statement, style="List Number", space_after=4)
    docx_para(document, p["closing"], space_after=18)
    for sline in p["signature"]:
        docx_para(document, sline or "", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    document.save(str(path))


if __name__ == "__main__":
    proposal_pdf = OUT / "GTNIC2026_WebDev_NamaTim.pdf"
    build_pdf(proposal_pdf, pdf_cover() + pdf_blocks(content.BLOCKS))
    build_docx(OUT / "GTNIC2026_WebDev_NamaTim.docx")

    build_pdf(OUT / "Surat_Pernyataan_Keaslian_Karya.pdf",
              [NextPageTemplate("body")] + pledge_blocks())
    build_pledge_docx(OUT / "Surat_Pernyataan_Keaslian_Karya.docx")

    for f in sorted(OUT.iterdir()):
        print(f"{f.name:44} {f.stat().st_size // 1024:>5} KB")
